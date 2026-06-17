# -*- coding: utf-8 -*-
"""ピエロ大好き人間 × Steenz 指づくりWS
イベント運営ワークフロー（準備〜事後）マニュアル / ダーク調・図解
"""
from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BG       = RGBColor(0x0B, 0x0B, 0x0D)
PANEL    = RGBColor(0x16, 0x16, 0x1A)
PANEL2   = RGBColor(0x1E, 0x1E, 0x23)
BLOOD    = RGBColor(0xB7, 0x12, 0x1C)
BLOOD_HI = RGBColor(0xE6, 0x1A, 0x1A)
INK      = RGBColor(0xEC, 0xEC, 0xEC)
MUTE     = RGBColor(0x9A, 0x9A, 0x9E)
LINEC    = RGBColor(0x40, 0x15, 0x18)
JP = "Yu Gothic UI"
W, H = 13.333, 7.5


def ea(run, name=JP):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)


def bg(slide, color=BG):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, In(W), In(H))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    s.shadow.inherit = False
    sp = s._element; sp.getparent().remove(sp); slide.shapes._spTree.insert(2, sp)


def rrect(slide, x, y, w, h, fill, line=None, radius=0.09, lw=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, In(x), In(y), In(w), In(h))
    try: s.adjustments[0] = radius
    except Exception: pass
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def rect(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, In(x), In(y), In(w), In(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill; s.line.fill.background()
    s.shadow.inherit = False
    return s


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    t = slide.shapes.add_textbox(In(x), In(y), In(w), In(h)).text_frame
    t.word_wrap = True; t.vertical_anchor = anchor
    t.margin_left = 0; t.margin_right = 0; t.margin_top = 0; t.margin_bottom = 0
    return t


def line(tf, runs, first=False, align=PP_ALIGN.LEFT, sa=3, ls=1.05):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(sa); p.space_before = Pt(0); p.line_spacing = ls
    for txt, sz, col, bold in runs:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = bold; r.font.color.rgb = col; ea(r)
    return p


def one(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, first=True, sa=3, ls=1.05):
    return line(tf, [(text, size, color, bold)], first=first, align=align, sa=sa, ls=ls)


prs = Presentation(); prs.slide_width = In(W); prs.slide_height = In(H)
BLANK = prs.slide_layouts[6]
PAGES = 10


def base(no, title, caption=None, danger=False):
    s = prs.slides.add_slide(BLANK); bg(s)
    rect(s, 0, 0, W, 1.4, PANEL); rect(s, 0, 1.4, W, 0.045, BLOOD)
    t = tb(s, 0.55, 0.15, 2.0, 1.1, MSO_ANCHOR.MIDDLE)
    one(t, no, 30, BLOOD, bold=True)
    t = tb(s, 2.45, 0.15, 10.3, 1.1, MSO_ANCHOR.MIDDLE)
    one(t, ("⚠ " if danger else "") + title, 26, INK, bold=True)
    top = 1.72
    if caption:
        c = tb(s, 0.95, 1.55, 11.6, 0.4); one(c, caption, 13, MUTE); top = 2.05
    f = tb(s, 0.95, 7.05, 9.0, 0.35, MSO_ANCHOR.MIDDLE)
    one(f, "指づくりWS 運営ワークフロー（準備〜事後）", 9, MUTE)
    return s, top


def pageno(s, idx):
    p = tb(s, 11.4, 7.05, 1.4, 0.35, MSO_ANCHOR.MIDDLE)
    one(p, f"{idx}/{PAGES-1}", 9, MUTE, align=PP_ALIGN.RIGHT)


def cols(n, gap=0.3, left=0.95, right=0.95):
    total = W - left - right
    cw = (total - gap * (n - 1)) / n
    return [left + i * (cw + gap) for i in range(n)], cw


def chead(s, x, y, w, label):
    h = tb(s, x + 0.25, y + 0.18, w - 0.5, 0.5)
    one(h, label, 16, BLOOD_HI, bold=True)


# ============================================================
# 表紙
# ============================================================
s = prs.slides.add_slide(BLANK); bg(s)
rect(s, 0, 0, 0.35, H, BLOOD)
rect(s, 0.95, 1.75, 8.0, 0.07, BLOOD)
t = tb(s, 1.0, 2.0, 11.6, 2.4)
one(t, "イベント運営", 44, INK, bold=True, sa=2)
one(t, "ワークフロー・マニュアル", 44, INK, bold=True, first=False)
t = tb(s, 1.0, 4.35, 11.6, 0.6)
one(t, "企画 → 収支 → 承認 → 準備 → 当日 → 事後処理", 20, BLOOD_HI, bold=True)
t = tb(s, 1.0, 5.1, 11.6, 1.6)
for i, ln in enumerate(["運営メンバー向け（イベントの立ち上げ〜締めまで）",
                         "「上から順にこなせば1イベントが回る」全体像",
                         "当日の運営手順は別冊『当日マニュアル』を参照"]):
    one(t, "▸  " + ln, 15, MUTE, first=(i == 0), sa=5)
rect(s, 0, 6.95, W, 0.55, PANEL)
t = tb(s, 1.0, 6.95, 11.6, 0.55, MSO_ANCHOR.MIDDLE)
line(t, [("⚠  ", 14, BLOOD_HI, True), ("③の承認が出るまで、応募・発注には進まない", 14, INK, True)], first=True)

# ============================================================
# 01 全体5ステップ
# ============================================================
s, top = base("01", "全体の流れ（5ステップ）", caption="運営メンバー向け。上から順にこなせばイベントは回る")
steps = [
    ("①", "イベント選定", "相性を見て\n企画書を作る"),
    ("②", "収支計画", "収益シートで\n見積り"),
    ("③", "承認", "渡邊さんに\nOKをもらう"),
    ("④", "準備・実施", "材料・人・告知\n→ 当日"),
    ("⑤", "事後処理", "給料申請・\n会計報告"),
]
xs, cw = cols(5, gap=0.25)
chh = 3.5
for i, (x, (nm, ttl, note)) in enumerate(zip(xs, steps)):
    y = top + 0.3
    rrect(s, x, y, cw, chh, PANEL, line=LINEC)
    t = tb(s, x, y + 0.3, cw, 0.8, MSO_ANCHOR.MIDDLE)
    one(t, nm, 34, BLOOD_HI, bold=True, align=PP_ALIGN.CENTER)
    t = tb(s, x + 0.1, y + 1.15, cw - 0.2, 0.8, MSO_ANCHOR.MIDDLE)
    one(t, ttl, 15, INK, bold=True, align=PP_ALIGN.CENTER, ls=1.0)
    t = tb(s, x + 0.1, y + 1.95, cw - 0.2, chh - 2.0)
    for j, nl in enumerate(note.split("\n")):
        one(t, nl, 12, MUTE, first=(j == 0), align=PP_ALIGN.CENTER, sa=2)
    if i < 4:
        a = tb(s, x + cw - 0.05, y, 0.35, chh, MSO_ANCHOR.MIDDLE)
        one(a, "›", 24, BLOOD_HI, bold=True, align=PP_ALIGN.CENTER)
band = top + 0.3 + chh + 0.2
rrect(s, 0.95, band, W - 1.9, 0.7, BLOOD, radius=0.12)
t = tb(s, 0.95, band, W - 1.9, 0.7, MSO_ANCHOR.MIDDLE)
one(t, "承認（③）が出るまで応募・発注に進まない／要修正なら②へ戻る", 15, INK, bold=True, align=PP_ALIGN.CENTER)
pageno(s, 1)

# ============================================================
# 02 STEP① 選定
# ============================================================
s, top = base("02", "STEP① イベントの選定", caption="出るイベントを探し、ピエロ大好き人間と相性が良いか判断する")
cards = [
    ("相性を見極める", ["判断材料：客層・過去の写真・出展者", "SNS・公式サイトを確認、可能なら現地へ行く", "主力＝デザフェス／ニコニコ超会議。新規も定期検討"]),
    ("企画書をつくる", ["相性OKなら企画書テンプレをコピーして記入", "開催日・場所・サイト・「今回挑戦すること」", "毎回同じだと飽きられる → 変化・トライを毎回決める"]),
]
yy = top + 0.15
for ttl, lns in cards:
    rrect(s, 0.95, yy, W - 1.9, 1.95, PANEL, line=LINEC)
    chead(s, 0.95, yy, W - 1.9, ttl)
    t = tb(s, 1.3, yy + 0.75, W - 2.6, 1.1)
    for j, ln in enumerate(lns):
        line(t, [("・ ", 14, BLOOD_HI, True), (ln, 15, INK, False)], first=(j == 0), sa=5)
    yy += 2.15
rrect(s, 0.95, yy, W - 1.9, 0.6, PANEL2, line=BLOOD)
t = tb(s, 0.95, yy, W - 1.9, 0.6, MSO_ANCHOR.MIDDLE)
one(t, "相性OK なら STEP② へ", 15, BLOOD_HI, bold=True, align=PP_ALIGN.CENTER)
pageno(s, 2)

# ============================================================
# 03 STEP② 収支計画
# ============================================================
s, top = base("03", "STEP② 収支計画をつくる", caption="会計シートで「いくら売れて、いくらかかるか」を見積もり、可否を決める")
items = [
    ("1", "収益シートをコピー", "テンプレを staff@steenz.jp でログインしてコピー（直接書き込まない）"),
    ("2", "席数を確定 → レイアウト", "席数6席。Canvaでブースレイアウト。什器は主催が用意する場合あり→要確認"),
    ("3", "売上−コストで可否判断", "客数・席数から想定売上、コストを算出。損益分岐点で実施可否を意思決定"),
]
yy = top + 0.15
for num, ttl, sub in items:
    rrect(s, 0.95, yy, W - 1.9, 1.3, PANEL, line=LINEC)
    t = tb(s, 1.2, yy, 1.0, 1.3, MSO_ANCHOR.MIDDLE)
    one(t, num, 32, BLOOD_HI, bold=True, align=PP_ALIGN.CENTER)
    t = tb(s, 2.3, yy + 0.22, W - 2.3 - 1.2, 0.95, MSO_ANCHOR.MIDDLE)
    one(t, ttl, 17, INK, bold=True, sa=3)
    one(t, sub, 13, MUTE, first=False)
    yy += 1.45
rrect(s, 0.95, yy, W - 1.9, 0.55, PANEL2, line=BLOOD)
t = tb(s, 0.95, yy, W - 1.9, 0.55, MSO_ANCHOR.MIDDLE)
one(t, "計画ができたら STEP③ へ", 14, BLOOD_HI, bold=True, align=PP_ALIGN.CENTER)
pageno(s, 3)

# ============================================================
# 04 STEP③ 承認
# ============================================================
s, top = base("04", "STEP③ 渡邊さんに承認をもらう", caption="計画をまとめて確認を依頼。共有に必ず含める4点セット")
quad = [("最大売上", "稼働率100％の売上"), ("最大コスト", "かかりうる最大費用"),
        ("最低限のコスト", "固定費（出展料など）"), ("品目 × 数量", "どの商品を何個売るか")]
xs, cw = cols(2, gap=0.5)
rh = 1.5; gy = 0.25
for i, (ttl, sub) in enumerate(quad):
    x = xs[i % 2]; y = top + 0.2 + (i // 2) * (rh + gy)
    rrect(s, x, y, cw, rh, PANEL, line=LINEC)
    t = tb(s, x + 0.35, y + 0.25, cw - 0.6, rh - 0.5, MSO_ANCHOR.MIDDLE)
    one(t, ttl, 22, BLOOD_HI, bold=True, sa=3)
    one(t, sub, 13, MUTE, first=False)
by = top + 0.2 + 2 * rh + gy + 0.15
rrect(s, 0.95, by, W - 1.9, 0.7, BLOOD, radius=0.12)
t = tb(s, 0.95, by, W - 1.9, 0.7, MSO_ANCHOR.MIDDLE)
one(t, "承認が出るまで応募・発注に進まない（要修正なら STEP② へ戻る）", 15, INK, bold=True, align=PP_ALIGN.CENTER)
pageno(s, 4)

# ============================================================
# 05 STEP④ 準備
# ============================================================
s, top = base("05", "STEP④ 準備して当日を実施", caption="承認された会計シートをもとに、材料・人・告知を準備する")
prep = [
    ("4-1", "材料の調達", "在庫SSで確認し発注。デザフェス備品はカタログ発注。領収書は会期月末までにメール依頼（宛名指定・再発行不可）"),
    ("4-2", "スタッフの収集", "声がけ時に交通費込み日給を伝える。LINEはイベントごとに1つ（渡邊さんを含める）"),
    ("4-3", "シフト作成", "自動生成テンプレを使用。1日目は1.5〜2h前/2日目1h前集合。休憩1人1h・30分ずらし"),
    ("4-4", "告知・アンケ・PR", "出展申込・整理券・SNS。アンケートは共通1本。PR・プロモは平沢くんに一任"),
    ("4-5", "当日の実施", "親ページを複製し【 】を埋めて配布 → 別冊『当日マニュアル』に沿って運営"),
]
xs, cw = cols(2, gap=0.5)
rh = 1.35; gy = 0.18
for i, (num, ttl, sub) in enumerate(prep):
    x = xs[i % 2]; y = top + 0.05 + (i // 2) * (rh + gy)
    rrect(s, x, y, cw, rh, PANEL, line=LINEC)
    t = tb(s, x + 0.25, y + 0.18, cw - 0.5, 0.5)
    line(t, [(num + "  ", 15, BLOOD_HI, True), (ttl, 16, INK, True)], first=True)
    t = tb(s, x + 0.25, y + 0.62, cw - 0.5, rh - 0.65)
    one(t, sub, 11.5, MUTE, ls=1.05)
pageno(s, 5)

# ============================================================
# 06 STEP⑤ 事後処理
# ============================================================
s, top = base("06", "STEP⑤ 事後処理", caption="表を上から消化。給料申請とお店側への会計報告が完了でイベント終了")
rows = [
    ("当日夜", "SNSお礼投稿（写真選別＋次回予告）", False),
    ("翌日", "売上・支出の最終集計を会計シートへ", False),
    ("翌日", "在庫棚卸し（必ず実施・補充把握）", True),
    ("会期月末まで", "レンタル備品の領収書をメール依頼（宛名指定・再発行不可）", True),
    ("+1週間", "振り返り会（KPT）→ 次回応募までに完了", False),
    ("月末まで", "スタッフへの給料支払い申請（厳守）", True),
    ("最後", "渡邊さんへ会計報告 → 完了でイベント終了 🎉", True),
]
tx = 0.95; tw = W - 1.9; c1 = 3.0; rh = 0.6; y = top + 0.1
rrect(s, tx, y, tw, rh, BLOOD, radius=0.06)
t = tb(s, tx + 0.25, y, c1, rh, MSO_ANCHOR.MIDDLE); one(t, "時期", 14, INK, bold=True)
t = tb(s, tx + c1 + 0.25, y, tw - c1 - 0.4, rh, MSO_ANCHOR.MIDDLE); one(t, "タスク", 14, INK, bold=True)
y += rh + 0.06
for tm, tsk, em in rows:
    rrect(s, tx, y, tw, rh, PANEL2 if em else PANEL, line=BLOOD_HI if em else LINEC, lw=1.4 if em else 1.0)
    t = tb(s, tx + 0.25, y, c1, rh, MSO_ANCHOR.MIDDLE)
    one(t, tm, 13, BLOOD_HI if em else INK, bold=True)
    t = tb(s, tx + c1 + 0.25, y, tw - c1 - 0.4, rh, MSO_ANCHOR.MIDDLE)
    one(t, tsk, 13, INK)
    y += rh + 0.06
pageno(s, 6)

# ============================================================
# 07 収益シートの使い方
# ============================================================
s, top = base("07", "収益シート（会計シート）の使い方", caption="収支計画（②）から会計報告（⑤）まで、このシート1枚で管理")
ph = [("イベント前", ["売上見立て・コスト", "損益分岐点を算出", "→ 承認用の数字づくり"]),
      ("当日", ["商品別の販売個数", "現金カウント（金種別）", "PayPayも記録"]),
      ("イベント後", ["売上突合（差額は理由記入）", "分配を自動計算", "→ 会計報告"])]
xs, cw = cols(3, gap=0.4)
for x, (ttl, lns) in zip(xs, ph):
    rrect(s, x, top + 0.1, cw, 2.1, PANEL, line=LINEC)
    rrect(s, x, top + 0.1, cw, 0.55, BLOOD, radius=0.14)
    t = tb(s, x, top + 0.13, cw, 0.5, MSO_ANCHOR.MIDDLE)
    one(t, ttl, 15, INK, bold=True, align=PP_ALIGN.CENTER)
    t = tb(s, x + 0.3, top + 0.8, cw - 0.55, 1.3)
    for j, ln in enumerate(lns):
        line(t, [("・ ", 12, BLOOD_HI, True), (ln, 12.5, INK, False)], first=(j == 0), sa=4)
ry = top + 2.45
strips = [
    ("人件費ルール", "Steenzスタッフ派遣分は記載しない／平澤くんの後輩・友人やSteenzスタッフの友人は記載OK（勝手に変えない）"),
    ("分配率", "ワークショップ・ガチャ＝Steenz 50％／平澤くん 50％　｜　作品系（手・足等）＝Steenz 30％／平澤くん 70％"),
    ("鉄則", "コピーは必ず staff@steenz.jp で／数式セルは触らない／現金は最初と最後に2人で確認・記録"),
]
for h, b in strips:
    rrect(s, 0.95, ry, W - 1.9, 0.62, PANEL2, line=BLOOD)
    t = tb(s, 1.2, ry, W - 2.2, 0.62, MSO_ANCHOR.MIDDLE)
    line(t, [(h + "：", 13, BLOOD_HI, True), (b, 12.5, INK, False)], first=True)
    ry += 0.72
pageno(s, 7)

# ============================================================
# 08 進捗管理・タスク表
# ============================================================
s, top = base("08", "進捗管理（タスク表で見失わない）", caption="頭で覚えず、フォルダ分け（親子構造）のタスク表に任せる")
folders = [
    ("💰 会計", ["見積もり・承認", "集計・支払い", "会計報告"]),
    ("🎪 イベント", ["選定・企画", "レイアウト・申込", "発注・当日・在庫"]),
    ("🧑‍🤝‍🧑 スタッフ", ["スタッフ集め", "LINE・シフト表", "1週間前の案内"]),
]
xs, cw = cols(3, gap=0.4)
for x, (ttl, lns) in zip(xs, folders):
    rrect(s, x, top + 0.1, cw, 2.4, PANEL, line=LINEC)
    rrect(s, x, top + 0.1, cw, 0.6, BLOOD, radius=0.12)
    t = tb(s, x, top + 0.13, cw, 0.55, MSO_ANCHOR.MIDDLE)
    one(t, ttl, 16, INK, bold=True, align=PP_ALIGN.CENTER)
    t = tb(s, x + 0.3, top + 0.9, cw - 0.55, 1.5)
    for j, ln in enumerate(lns):
        line(t, [("▸ ", 13, BLOOD_HI, True), (ln, 14, INK, False)], first=(j == 0), sa=7)
by = top + 2.7
rrect(s, 0.95, by, W - 1.9, 1.35, PANEL2, line=LINEC)
chead(s, 0.95, by, W - 1.9, "使い方")
t = tb(s, 1.3, by + 0.7, W - 2.6, 0.6)
one(t, "① テンプレを複製しタイトルをイベント名に変更　② 進捗を 未着手→進行中→完了 で更新", 13.5, INK, sa=4)
one(t, "③ 中断するときは進行中タスクのメモに「★次やること」を一言残す", 13.5, INK, first=False)
pageno(s, 8)

# ============================================================
# 09 テンプレ・リンク集
# ============================================================
s, top = base("09", "テンプレ・リンク集", caption="イベントごとにコピーして使う主要ツール")
groups = [
    ("📄 テンプレ", ["企画書テンプレート", "会計報告テンプレート", "イベント準備タスク表", "支払い依頼シート"]),
    ("📊 シート類", ["収益シート（staff@steenzでコピー）", "シフト表（自動生成）", "備品リスト・在庫管理SS", "これまでの収益シート（実績）"]),
    ("🎨 Canva素材", ["ブースレイアウト／完成イメージ", "整理券テンプレ①②", "看板デザイン①②③", "アンケートQRコード"]),
]
xs, cw = cols(3, gap=0.4)
for x, (ttl, lns) in zip(xs, groups):
    rrect(s, x, top + 0.1, cw, 3.2, PANEL, line=LINEC)
    rrect(s, x, top + 0.1, cw, 0.6, BLOOD, radius=0.12)
    t = tb(s, x, top + 0.13, cw, 0.55, MSO_ANCHOR.MIDDLE)
    one(t, ttl, 15, INK, bold=True, align=PP_ALIGN.CENTER)
    t = tb(s, x + 0.3, top + 0.95, cw - 0.55, 2.2)
    for j, ln in enumerate(lns):
        line(t, [("・ ", 13, BLOOD_HI, True), (ln, 13, INK, False)], first=(j == 0), sa=9)
by = top + 3.45
rrect(s, 0.95, by, W - 1.9, 0.6, PANEL2, line=BLOOD)
t = tb(s, 0.95, by, W - 1.9, 0.6, MSO_ANCHOR.MIDDLE)
one(t, "⚠ 出展申込のアカウント情報（ID・PW）はこのマニュアルには載せない（運営メンバーのみで管理）", 12.5, INK, bold=True, align=PP_ALIGN.CENTER)
pageno(s, 9)

out = "指づくりWS_運営ワークフロー_ダーク.pptx"
prs.save(out)
print("saved:", out, "/ slides:", len(prs.slides._sldIdLst))
