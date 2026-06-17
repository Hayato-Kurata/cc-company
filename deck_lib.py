# -*- coding: utf-8 -*-
"""指づくりWS マニュアル共通ライブラリ（ダーク・ホラー調 / 図解）"""
from pptx import Presentation
from pptx.util import Inches as In, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- パレット ----
BG       = RGBColor(0x0B, 0x0B, 0x0D)
PANEL    = RGBColor(0x16, 0x16, 0x1A)
PANEL2   = RGBColor(0x1E, 0x1E, 0x23)
BLOOD    = RGBColor(0xB7, 0x12, 0x1C)
BLOOD_HI = RGBColor(0xE6, 0x1A, 0x1A)
INK      = RGBColor(0xEC, 0xEC, 0xEC)
MUTE     = RGBColor(0x9A, 0x9A, 0x9E)
LINEC    = RGBColor(0x40, 0x15, 0x18)
JP = "Yu Gothic UI"
W, H = 13.333, 7.5


def ea(run, name=JP):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)


def bg(slide, color=BG):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, In(W), In(H))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    s.shadow.inherit = False
    sp = s._element; sp.getparent().remove(sp); slide.shapes._spTree.insert(2, sp)


def rrect(slide, x, y, w, h, fill, line=None, radius=0.09, lw=1.0):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, In(x), In(y), In(w), In(h))
    try: s.adjustments[0] = radius
    except Exception: pass
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def rect(slide, x, y, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, In(x), In(y), In(w), In(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill; s.line.fill.background()
    s.shadow.inherit = False
    return s


def tb(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    t = slide.shapes.add_textbox(In(x), In(y), In(w), In(h)).text_frame
    t.word_wrap = True; t.vertical_anchor = anchor
    t.margin_left = 0; t.margin_right = 0; t.margin_top = 0; t.margin_bottom = 0
    return t


def line(tf, runs, first=False, align=PP_ALIGN.LEFT, sa=3, ls=1.05):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align; p.space_after = Pt(sa); p.space_before = Pt(0); p.line_spacing = ls
    for txt, sz, col, bold in runs:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = bold; r.font.color.rgb = col; ea(r)
    return p


def one(tf, text, size, color, bold=False, align=PP_ALIGN.LEFT, first=True, sa=3, ls=1.05):
    return line(tf, [(text, size, color, bold)], first=first, align=align, sa=sa, ls=ls)


def cols(n, gap=0.3, left=0.95, right=0.95):
    total = W - left - right
    cw = (total - gap * (n - 1)) / n
    return [left + i * (cw + gap) for i in range(n)], cw


def chead(s, x, y, w, label):
    h = tb(s, x + 0.25, y + 0.18, w - 0.5, 0.5)
    one(h, label, 16, BLOOD_HI, bold=True)


def new_prs():
    prs = Presentation(); prs.slide_width = In(W); prs.slide_height = In(H)
    return prs, prs.slide_layouts[6]


def cover(prs, blank, line1, line2, subtitle, bullets, foot):
    s = prs.slides.add_slide(blank); bg(s)
    rect(s, 0, 0, 0.35, H, BLOOD)
    rect(s, 0.95, 1.75, 8.0, 0.07, BLOOD)
    t = tb(s, 1.0, 2.0, 11.6, 2.4)
    one(t, line1, 44, INK, bold=True, sa=2)
    if line2:
        one(t, line2, 44, INK, bold=True, first=False)
    t = tb(s, 1.0, 4.35, 11.6, 0.6)
    one(t, subtitle, 20, BLOOD_HI, bold=True)
    t = tb(s, 1.0, 5.1, 11.6, 1.6)
    for i, ln in enumerate(bullets):
        one(t, "▸  " + ln, 15, MUTE, first=(i == 0), sa=5)
    rect(s, 0, 6.95, W, 0.55, PANEL)
    t = tb(s, 1.0, 6.95, 11.6, 0.55, MSO_ANCHOR.MIDDLE)
    line(t, [("⚠  ", 14, BLOOD_HI, True), (foot, 14, INK, True)], first=True)
    return s


def base(prs, blank, no, title, foot_label, caption=None, danger=False):
    s = prs.slides.add_slide(blank); bg(s)
    rect(s, 0, 0, W, 1.4, PANEL); rect(s, 0, 1.4, W, 0.045, BLOOD)
    t = tb(s, 0.55, 0.15, 1.7, 1.1, MSO_ANCHOR.MIDDLE)
    one(t, no, 32, BLOOD, bold=True)
    t = tb(s, 2.15, 0.15, 10.6, 1.1, MSO_ANCHOR.MIDDLE)
    one(t, ("⚠ " if danger else "") + title, 25, INK, bold=True, ls=1.0)
    top = 1.72
    if caption:
        c = tb(s, 0.95, 1.55, 11.6, 0.4); one(c, caption, 13, MUTE); top = 2.05
    f = tb(s, 0.95, 7.05, 9.0, 0.35, MSO_ANCHOR.MIDDLE)
    one(f, foot_label, 9, MUTE)
    return s, top


def pageno(s, idx, total):
    p = tb(s, 11.4, 7.05, 1.4, 0.35, MSO_ANCHOR.MIDDLE)
    one(p, f"{idx}/{total}", 9, MUTE, align=PP_ALIGN.RIGHT)
