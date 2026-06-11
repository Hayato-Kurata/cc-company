#!/usr/bin/env python3
"""小学館 経費精算シート 自動入力スクリプト

template.xlsx（固定情報入り・数式保持）を読み込み、JSON で渡した明細を
①経費請求明細 / ②交通費請求明細 シートに転記して完成版 Excel を出力する。
合計・消費税・請求金額はテンプレート内の数式が自動計算するため触らない。

使い方:
    python3 fill_expense.py input.json -o 2026-02_経費.xlsx

入力 JSON（すべて任意。空なら該当シートはそのまま）:
{
  "請求日": "2026-02-28",

  "経費": [
    {"日付":"2026-02-28","支払先":"セブンイレブン","摘要":"外務省収録(お茶)",
     "金額":1663,"税率":"8%","インボイス":true}
  ],

  "交通費": [
    {"日付":"2026-02-16","出社":true},                     // 中央林間⇔神保町 出社 1180円
    {"日付":"2026-02-28","訪問先":"エイトリンクス","摘要":"外務省収録スタッフ"}
  ],

  // さらに簡単に：出社日を並べるだけで交通費行を一括生成
  "出社日": ["2026-02-16","2026-02-18","2026-02-20","2026-02-26","2026-02-27"]
}

税率の指定: "10%" / "8%"（軽減）/ "非課税" のいずれか。省略時は 10%。
"""
import argparse
import datetime
import json
import sys
from pathlib import Path

import openpyxl

TEMPLATE = Path(__file__).with_name("template.xlsx")
DETAIL_FIRST_ROW = 16
DETAIL_LAST_ROW = 35  # 20 行（16〜35）。超える場合は分割が必要

# 通勤（出社）の既定値。ほぼ毎回これなので一言で入る
COMMUTE_DEFAULTS = {
    "手段": "電車",
    "区間from": "中央林間",
    "区間sep": "⇔",
    "区間to": "神保町",
    "訪問先": "小学館ビル",
    "摘要": "出社",
    "金額": 1180,
    "インボイス": True,
}


def parse_date(v):
    if v in (None, ""):
        return None
    if isinstance(v, datetime.datetime):
        return v
    if isinstance(v, datetime.date):
        return datetime.datetime(v.year, v.month, v.day)
    s = str(v).strip().replace("/", "-")
    return datetime.datetime.strptime(s, "%Y-%m-%d")


def tax_cell_value(rate):
    """税率指定 -> セルに書く値（既存ファイルの仕様に合わせる）。
    10% は数値 0.1（パーセント書式）、8% は '軽減8%'、非課税は '非/不課税'。"""
    if rate in (None, "", "10%", "10", 10, 0.1):
        return 0.1
    s = str(rate)
    if "8" in s or "軽減" in s:
        return "軽減8%"
    if "非" in s or "不課税" in s:
        return "非/不課税"
    return 0.1


def fill_expenses(ws, rows):
    if len(rows) > (DETAIL_LAST_ROW - DETAIL_FIRST_ROW + 1):
        raise ValueError(
            f"経費明細が{len(rows)}件あります。1枚あたり最大20件です。分割してください。"
        )
    for i, item in enumerate(rows):
        r = DETAIL_FIRST_ROW + i
        ws.cell(row=r, column=2).value = parse_date(item.get("日付"))      # B 利用日
        ws.cell(row=r, column=3).value = item.get("支払先", "")            # C 支払先
        ws.cell(row=r, column=4).value = item.get("摘要", "")              # D 摘要
        ws.cell(row=r, column=7).value = item.get("金額")                  # G 金額
        ws.cell(row=r, column=9).value = tax_cell_value(item.get("税率"))  # I 税率
        ws.cell(row=r, column=11).value = i + 1                            # K 証票番号
        ws.cell(row=r, column=13).value = bool(item.get("インボイス", True))  # M インボイス


def fill_transit(ws, rows):
    if len(rows) > (DETAIL_LAST_ROW - DETAIL_FIRST_ROW + 1):
        raise ValueError(
            f"交通費明細が{len(rows)}件あります。1枚あたり最大20件です。分割してください。"
        )
    for i, item in enumerate(rows):
        r = DETAIL_FIRST_ROW + i
        d = dict(COMMUTE_DEFAULTS)  # 既定（出社パターン）から開始し、指定分で上書き
        d.update({k: v for k, v in item.items() if k != "出社"})
        ws.cell(row=r, column=2).value = parse_date(item.get("日付"))   # B 利用日
        ws.cell(row=r, column=3).value = d["手段"]                       # C 手段
        ws.cell(row=r, column=4).value = d["区間from"]                   # D 区間(発)
        ws.cell(row=r, column=5).value = d["区間sep"]                    # E ⇔/→
        ws.cell(row=r, column=6).value = d["区間to"]                     # F 区間(着)
        ws.cell(row=r, column=7).value = d["訪問先"]                     # G 訪問先
        ws.cell(row=r, column=8).value = d["摘要"]                       # H 摘要
        ws.cell(row=r, column=10).value = d["金額"]                      # J 金額
        ws.cell(row=r, column=13).value = i + 1                          # M 証票番号
        ws.cell(row=r, column=15).value = bool(d["インボイス"])          # O インボイス


def build(data, template=TEMPLATE):
    wb = openpyxl.load_workbook(template)  # 数式を保持

    seikyuubi = parse_date(data.get("請求日"))

    keihi = data.get("経費", [])
    if keihi:
        ws = wb["①経費請求明細"]
        if seikyuubi:
            ws["I1"] = seikyuubi
        fill_expenses(ws, keihi)

    # 交通費：明示の "交通費" 行 ＋ "出社日" のショートカットを結合
    transit = list(data.get("交通費", []))
    for d in data.get("出社日", []):
        transit.append({"日付": d, "出社": True})
    # 日付順に並べる
    transit.sort(key=lambda x: str(x.get("日付", "")))
    if transit:
        ws2 = wb["②交通費請求明細"]
        if seikyuubi:
            ws2["L1"] = seikyuubi
        fill_transit(ws2, transit)

    return wb


def main(argv=None):
    p = argparse.ArgumentParser(description="小学館 経費精算シート 自動入力")
    p.add_argument("input", help="明細を記述した JSON ファイル")
    p.add_argument("-o", "--output", help="出力 xlsx パス", default="経費精算.xlsx")
    p.add_argument("-t", "--template", default=str(TEMPLATE), help="テンプレート xlsx")
    args = p.parse_args(argv)

    data = json.loads(Path(args.input).read_text(encoding="utf-8"))
    wb = build(data, template=Path(args.template))
    wb.save(args.output)

    n_k = len(data.get("経費", []))
    n_t = len(data.get("交通費", [])) + len(data.get("出社日", []))
    print(f"✓ 出力: {args.output}（経費 {n_k} 件 / 交通費 {n_t} 件）")
    print("  合計・消費税・請求金額は Excel で開くと数式が自動計算します。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
